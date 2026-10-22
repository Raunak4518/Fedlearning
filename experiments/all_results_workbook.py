"""Every experiment and every measured metric in one Excel workbook.

    python experiments/all_results_workbook.py   -> docs/paper_assets/all_results.xlsx

Sheets
  README            what is in the workbook and how to read it
  Experiments       one row per result folder (spec title, notebook, runs, datasets, seeds) and every planned
                    Kaggle job that never returned results
  All runs          one row per logged run (every runs.jsonl in experiments/results), one column per scalar metric;
                    blank = not measured in that run
  Summary           one row per setting (experiment, method, dataset, IF, alpha, K, ...): number of runs and the
                    mean (and SD) over seeds of the main metrics, as AVERAGEIF formulas over "All runs"; plus the
                    one experiment whose per-run file never came back (K01 CIFAR-10), from its summary table
  Per-class recall  per-class recall / precision lists of every run and readout (class 0 = most common)
  Per-class other   other per-class lists (generator fidelity, row norms, correction scales, ...)
  MovieLens E40     the federated matrix-factorisation runs (different metrics)
  Other results     every scalar in result files that are not runs.jsonl (D01, T01, T02, docs/research/math_checks)
  Metric glossary   what each column prefix and suffix means
Re-run after adding results.
"""
import glob
import json
import os
import re
from collections import OrderedDict, defaultdict

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, ".."))
RES = os.path.join(HERE, "results")
MC = os.path.join(REPO, "docs", "research", "math_checks")
OUT = os.path.join(REPO, "docs", "paper_assets", "all_results.xlsx")

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1C4E9A")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=14, bold=True)

ID_COLS = ["experiment", "result_folder", "notebook", "run_id", "label", "dataset", "IF", "alpha", "K", "seed",
           "pfrac", "tail_perm", "quick", "gate_branch", "code_commit", "repeat_of_earlier_row"]
READOUTS = ["final", "er", "er_own", "ccvr", "ccvr_own", "fed3r", "fed3r_own", "ncm", "ncm_own", "bbc", "oracle"]
SUFFIX = ["bal", "acc", "f1", "precision", "gmean", "worst", "worst_arch", "head", "medium", "tail", "rec_std",
          "ece", "nll", "brier", "local", "ens_bal", "ens_acc"]
FRACTION_SUFFIX = ("_bal", "_acc", "_f1", "_precision", "_gmean", "_worst", "_worst_arch", "_head", "_medium", "_tail",
                   "_rec_std", "_ece", "_local", "_ens_bal", "_ens_acc")
SUMMARY_METRICS = ["final_bal", "final_acc", "final_f1", "final_precision", "final_gmean", "final_worst", "final_tail",
                   "best_mean_acc", "er_bal", "er_f1", "er_gmean", "er_worst", "er_tail", "ccvr_bal", "fed3r_bal",
                   "ncm_bal", "bbc_bal", "oracle_bal"]
SD_METRICS = ["final_bal", "er_bal"]


EXTRA_TITLES = {  # experiments run by standalone scripts, not specs
    "D01_component_evidence": "Component evidence: generator fidelity, row norms and stage-III metrics per component (CPU)",
    "D01_kaggle": "Component evidence: generator fidelity, row norms and stage-III metrics per component (Kaggle)",
    "E40_fedembed_mf": "Rare-item collapse in federated matrix factorisation on MovieLens-1M",
    "T01_collapse_threshold": "Instrumented generator: class-row norms per client and round (collapse diagnosis)",
    "T02_dilution_rates": "Learning rate of rare class rows under plain averaging, without decay",
    "T03_exact_readout": "Exact readout from aggregated sums vs pooled-real refit (validation of ER)",
}


def spec_titles():
    out = dict(EXTRA_TITLES)
    for f in glob.glob(os.path.join(HERE, "specs", "*.py")):
        m = re.search(r'title="([^"]*)"', open(f, encoding="utf-8").read())
        out[os.path.basename(f)[:-3]] = m.group(1) if m else ""
    return out


def ledger_jobs():
    """job rows of LEDGER.md: notebook, args, result folder, planned, done, status."""
    rows = []
    for line in open(os.path.join(HERE, "LEDGER.md"), encoding="utf-8"):
        if line.startswith("| NB") or line.startswith("| local"):
            c = [x.strip() for x in line.strip().strip("|").split("|")]
            if len(c) >= 6:
                rows.append(dict(job=c[0], args=c[1], folder=c[2], planned=c[3], done=c[4], status=c[5]))
    return rows


def spec_of(folder):
    s = folder.split("__")[0]
    return s[:-6] if s.endswith("_quick") else s


def is_fraction(k):
    return k.endswith(FRACTION_SUFFIX) or k in ("best_mean_acc", "fidelity_head", "fidelity_tail", "referee_test_acc",
                                               "ncm_acc", "oracle_acc")


def load_runs():
    runs, e40 = [], []
    for d in sorted(glob.glob(os.path.join(RES, "*", "runs.jsonl"))):
        folder = os.path.basename(os.path.dirname(d))
        for line in open(d, encoding="utf-8"):
            if not line.strip():
                continue
            r = json.loads(line)
            r["_folder"] = folder
            (runs if "label" in r else e40).append(r)
    return runs, e40


def flat_scalars(obj, prefix=""):
    """(key, value) for every number / bool / short string in a nested JSON object; lists of numbers as text."""
    out = []
    if isinstance(obj, dict):
        for k, v in obj.items():
            out += flat_scalars(v, f"{prefix}.{k}" if prefix else str(k))
    elif isinstance(obj, list):
        if obj and all(isinstance(x, (int, float)) for x in obj):
            if len(obj) <= 120:
                out.append((prefix, ", ".join(f"{x:.4g}" if isinstance(x, float) else str(x) for x in obj)))
            else:
                out.append((prefix, f"[{len(obj)} numbers]"))
        else:
            for i, v in enumerate(obj[:200]):
                out += flat_scalars(v, f"{prefix}[{i}]")
    elif isinstance(obj, (int, float, bool)) or (isinstance(obj, str) and len(obj) < 200):
        out.append((prefix, obj))
    return out


def style_header(ws, row=1):
    for c in ws[row]:
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(wrap_text=True, vertical="top")


def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def main():
    runs, e40 = load_runs()
    titles = spec_titles()
    jobs = ledger_jobs()
    nb_of = defaultdict(list)
    for j in jobs:
        nb_of[j["folder"]].append(j["job"])

    # ---- column order for "All runs"
    scalar_keys = OrderedDict()
    for r in runs:
        for k, v in r.items():
            if k.startswith("_") or k in ID_COLS or k == "method":
                continue
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                scalar_keys[k] = True
    ordered = []
    for ro in READOUTS:
        ordered += [f"{ro}_{s}" for s in SUFFIX if f"{ro}_{s}" in scalar_keys]
    ordered += [k for k in ("best_mean_acc",) if k in scalar_keys]
    ordered += sorted(k for k in scalar_keys if k.startswith("val_") and k not in ordered)
    ordered += sorted(k for k in scalar_keys if k not in ordered)

    wb = Workbook()
    readme = wb.active
    readme.title = "README"

    # ---- All runs
    ws = wb.create_sheet("All runs")
    header = ID_COLS + ["group_key"] + ordered
    ws.append(header)
    seen = {}
    groups = OrderedDict()
    for i, r in enumerate(runs, start=2):
        rid = r.get("run_id") or f"{r.get('label')}|{r.get('dataset')}|IF{r.get('IF')}|a{r.get('alpha')}|K{r.get('K')}|s{r.get('seed')}"
        exp = spec_of(r["_folder"])
        rep = seen.get((exp, rid))
        seen.setdefault((exp, rid), i)
        quick = bool(r.get("quick")) or r["_folder"].endswith("_quick") or r.get("quick_scale", 1.0) < 1.0
        gk = "|".join(str(x) for x in (exp, r.get("label"), r.get("dataset"), r.get("IF"), r.get("alpha"), r.get("K"),
                                        r.get("pfrac", ""), "tp" if r.get("tail_perm") else "", "quick" if quick else ""))
        groups.setdefault(gk, dict(exp=exp, label=r.get("label"), dataset=r.get("dataset"), IF=r.get("IF"),
                                   alpha=r.get("alpha"), K=r.get("K"), pfrac=r.get("pfrac"), tp=bool(r.get("tail_perm")),
                                   quick=quick, seeds=set(), folders=set()))
        groups[gk]["seeds"].add(r.get("seed"))
        groups[gk]["folders"].add(r["_folder"])
        row = [exp, r["_folder"], ", ".join(nb_of.get(r["_folder"], [])), rid, r.get("label"), r.get("dataset"),
               r.get("IF"), r.get("alpha"), r.get("K"), r.get("seed"), r.get("pfrac"), bool(r.get("tail_perm")) or None,
               quick or None, r.get("gate_branch"), r.get("code_commit"), rep, gk]
        row += [r.get(k) if isinstance(r.get(k), (int, float)) and not isinstance(r.get(k), bool) else None for k in ordered]
        ws.append(row)
    n_rows = len(runs) + 1
    style_header(ws)
    ws.freeze_panes = "G2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{n_rows}"
    for j, k in enumerate(header, 1):
        col = get_column_letter(j)
        ws.column_dimensions[col].width = 34 if k in ("run_id", "group_key") else 24 if k in ("result_folder", "label", "experiment", "notebook") else 12
        if k in ordered and is_fraction(k):
            for c in ws[col][1:]:
                c.number_format = "0.0%"
        elif k in ordered:
            for c in ws[col][1:]:
                c.number_format = "0.0000"
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = BODY
    col_of = {k: get_column_letter(j) for j, k in enumerate(header, 1)}

    # ---- Summary (formulas over All runs)
    ss = wb.create_sheet("Summary")
    sh = ["experiment", "method", "dataset", "IF", "alpha", "K", "participation", "random rare classes", "quick pass",
          "seeds", "result folders", "runs"] + [f"{m} (mean)" for m in SUMMARY_METRICS] + [f"{m} (SD)" for m in SD_METRICS]
    ss.append(sh)
    gcol = f"'All runs'!${col_of['group_key']}$2:${col_of['group_key']}${n_rows}"
    for i, (gk, g) in enumerate(groups.items(), start=2):
        seeds = ", ".join(str(s) for s in sorted(x for x in g["seeds"] if x is not None))
        row = [g["exp"], g["label"], g["dataset"], g["IF"], g["alpha"], g["K"], g["pfrac"], "yes" if g["tp"] else None,
               "yes" if g["quick"] else None, seeds, ", ".join(sorted(g["folders"])), f'=COUNTIF({gcol},"{gk}")']
        for m in SUMMARY_METRICS:
            if m not in col_of:
                row.append(None)
                continue
            rng = f"'All runs'!${col_of[m]}$2:${col_of[m]}${n_rows}"
            row.append(f'=IFERROR(AVERAGEIF({gcol},"{gk}",{rng}),"")')
        for m in SD_METRICS:
            rng = f"'All runs'!${col_of[m]}$2:${col_of[m]}${n_rows}"
            mean_cell = f"{get_column_letter(sh.index(m + ' (mean)') + 1)}{i}"
            row.append(f'=IFERROR(SQRT(SUMPRODUCT(({gcol}="{gk}")*({rng}<>"")*({rng}-{mean_cell})^2)'
                       f'/(COUNTIFS({gcol},"{gk}",{rng},"<>")-1)),"")')
        ss.append(row)
    # K01 CIFAR-10: only the summary table came back (3-seed means as printed on Kaggle)
    k01 = os.path.join(RES, "K01_kaggle_cifar10_svhn__cifar10", "summary.md")
    if os.path.exists(k01):
        setting = None
        for line in open(k01, encoding="utf-8"):
            m = re.match(r"## (\w+)\s+IF=([\d.]+)\s+alpha=(\S+)\s+K=(\d+)", line)
            if m:
                setting = m.groups()
                cols = None
                continue
            if line.startswith("| label"):
                cols = [x.strip() for x in line.strip().strip("|").split("|")]
                continue
            if setting and line.startswith("|") and not line.startswith("|---") and cols:
                c = [x.strip() for x in line.strip().strip("|").split("|")]
                vals = dict(zip(cols, c))
                mean = lambda s: float(s.split("±")[0]) / 100 if s and s[0].isdigit() else None
                row = ["K01_kaggle_cifar10_svhn", vals["label"], setting[0], float(setting[1]),
                       None if setting[2] == "None" else float(setting[2]), int(setting[3]), None, None, None,
                       "0, 1, 2", "K01_kaggle_cifar10_svhn__cifar10 (summary table only; per-run file not returned)",
                       int(vals["n"])]
                for mtr in SUMMARY_METRICS:
                    row.append(mean(vals.get(mtr, "")) if mtr in ("final_bal", "final_tail", "best_mean_acc") else None)
                row += [None] * len(SD_METRICS)
                ss.append(row)
    style_header(ss)
    ss.freeze_panes = "D2"
    ss.auto_filter.ref = f"A1:{get_column_letter(len(sh))}{ss.max_row}"
    set_widths(ss, [26, 30, 10, 7, 7, 6, 10, 10, 8, 12, 30, 6] + [11] * (len(SUMMARY_METRICS) + len(SD_METRICS)))
    for row in ss.iter_rows(min_row=2):
        for j, c in enumerate(row):
            c.font = BODY
            if j >= 12:
                c.number_format = "0.0%"

    # ---- Per-class recall / precision, and other per-class lists
    pc = wb.create_sheet("Per-class recall")
    po = wb.create_sheet("Per-class other")
    maxC = 100
    pc.append(["run_id", "result_folder", "dataset", "readout", "measure"] + [f"class {c}" for c in range(maxC)])
    po.append(["run_id", "result_folder", "dataset", "quantity"] + [f"[{c}]" for c in range(maxC)])
    for r in runs:
        rid = r.get("run_id") or r.get("label")
        for k, v in r.items():
            if not (isinstance(v, list) and v and all(isinstance(x, (int, float)) for x in v) and len(v) <= maxC):
                continue
            if k.endswith("_recall") or k.endswith("_precision_c"):
                ro, meas = (k[:-7], "recall") if k.endswith("_recall") else (k[:-12], "precision")
                pc.append([rid, r["_folder"], r.get("dataset"), ro, meas] + list(v))
            elif k not in ("client_counts",):
                po.append([rid, r["_folder"], r.get("dataset"), k] + list(v))
    for sh_ in (pc, po):
        style_header(sh_)
        sh_.freeze_panes = "F2" if sh_ is pc else "E2"
        sh_.auto_filter.ref = f"A1:{get_column_letter(sh_.max_column)}{sh_.max_row}"
        sh_.column_dimensions["A"].width = 46
        sh_.column_dimensions["B"].width = 30
        for row in sh_.iter_rows(min_row=2):
            for c in row:
                c.font = BODY
    for row in pc.iter_rows(min_row=2, min_col=6):
        for c in row:
            c.number_format = "0.0%"

    # ---- MovieLens E40
    me = wb.create_sheet("MovieLens E40")
    ekeys = []
    for r in e40:
        for k, v in r.items():
            if not k.startswith("_") and isinstance(v, (int, float, str)) and k not in ekeys:
                ekeys.append(k)
    me.append(["result_folder"] + ekeys)
    for r in e40:
        me.append([r["_folder"]] + [r.get(k) for k in ekeys])
    style_header(me)
    me.freeze_panes = "B2"
    for row in me.iter_rows(min_row=2):
        for c in row:
            c.font = BODY

    # ---- Other results (non-runs.jsonl files)
    ot = wb.create_sheet("Other results")
    ot.append(["source", "file", "record", "field", "value"])
    sources = []
    for d in ("D01_component_evidence", "D01_kaggle", "T01_collapse_threshold", "T02_dilution_rates"):
        for f in sorted(glob.glob(os.path.join(RES, d, "*.json"))):
            sources.append((d, f))
    for f in sorted(glob.glob(os.path.join(MC, "*.json")) + glob.glob(os.path.join(MC, "*.jsonl"))):
        sources.append(("math_checks (CPU)", f))
    for src, f in sources:
        try:
            if f.endswith(".jsonl"):
                recs = [json.loads(line) for line in open(f, encoding="utf-8") if line.strip()]
            else:
                recs = [json.load(open(f, encoding="utf-8"))]
        except Exception as e:  # noqa: BLE001
            ot.append([src, os.path.relpath(f, REPO), "", "unreadable", str(e)[:200]])
            continue
        for i, rec in enumerate(recs):
            for k, v in flat_scalars(rec):
                ot.append([src, os.path.relpath(f, REPO), i, k, v])
    style_header(ot)
    ot.freeze_panes = "A2"
    ot.auto_filter.ref = f"A1:E{ot.max_row}"
    set_widths(ot, [20, 56, 8, 60, 40])
    for row in ot.iter_rows(min_row=2):
        for c in row:
            c.font = BODY

    # ---- Experiments index
    ex = wb.create_sheet("Experiments", 1)
    ex.append(["experiment", "title (from the spec)", "result folder", "notebook / job", "runs logged", "datasets",
               "methods", "seeds", "status"])
    by_folder = defaultdict(list)
    for r in runs:
        by_folder[r["_folder"]].append(r)
    for r in e40:
        by_folder[r["_folder"]].append(r)
    folders = sorted(set(os.path.basename(p.rstrip("/")) for p in glob.glob(os.path.join(RES, "*/"))))
    for folder in folders:
        rs = by_folder.get(folder, [])
        exp = spec_of(folder)
        status = next((j["status"] for j in jobs if j["folder"] == folder), "")
        if not rs:
            status = status or "results in summary / diagnostic files only (see Summary or Other results)"
        ex.append([exp, titles.get(exp, ""), folder, ", ".join(nb_of.get(folder, [])), len(rs),
                   ", ".join(sorted({str(r.get("dataset", r.get("task", ""))) for r in rs})),
                   len({r.get("label", r.get("arm")) for r in rs}),
                   ", ".join(str(s) for s in sorted({r.get("seed") for r in rs if r.get("seed") is not None})),
                   status[:300]])
    for j in jobs:
        if j["folder"] not in folders:
            ex.append([spec_of(j["folder"]), titles.get(spec_of(j["folder"]), ""), j["folder"], j["job"], 0, "", "", "",
                       "planned; no results returned (" + j["status"][:200] + ")"])
    style_header(ex)
    ex.freeze_panes = "B2"
    ex.auto_filter.ref = f"A1:I{ex.max_row}"
    set_widths(ex, [30, 60, 36, 30, 8, 22, 8, 12, 60])
    for row in ex.iter_rows(min_row=2):
        for c in row:
            c.font = BODY
            c.alignment = Alignment(wrap_text=True, vertical="top")

    # ---- Metric glossary
    gl = wb.create_sheet("Metric glossary")
    gl.append(["column part", "meaning"])
    gloss = [
        ("final_", "the trained heads as they are at the last round (no readout fix)"),
        ("er_", "exact readout: each architecture's last layer replaced by the class-balanced LDA computed from sums over all clients"),
        ("er_own_", "exact readout from the architecture's own clients only"),
        ("ccvr_ / ccvr_own_", "CCVR readout (classifier re-trained on virtual features from federated class statistics), all / own clients"),
        ("fed3r_ / fed3r_own_", "Fed3R closed-form ridge readout, all / own clients"),
        ("ncm_ / ncm_own_", "nearest-class-mean readout (FedNCM), all / own clients"),
        ("bbc_", "balanced bias calibration of the heads on synthetic data"),
        ("oracle_", "diagnostic only: last layer re-fitted on pooled real data (not federated)"),
        ("val_", "the same metric on the 2,000-image validation split (used to choose settings)"),
        ("_bal", "balanced accuracy = mean over classes of recall (macro recall)"),
        ("_acc", "plain accuracy"),
        ("_f1", "macro-F1 (mean over classes of per-class F1)"),
        ("_precision", "macro precision (mean over classes; logged only by runs from 22 Oct 2026)"),
        ("_gmean", "geometric mean of class recalls (each floored at 1%)"),
        ("_worst", "recall of the worst class"),
        ("_worst_arch", "balanced accuracy of the weakest architecture"),
        ("_head / _medium / _tail", "mean recall of the most common 30% / middle 30% / rarest 40% of classes"),
        ("_rec_std", "spread (SD) of class recalls; lower = fairer across classes"),
        ("_ece / _nll / _brier", "calibration: expected calibration error, negative log-likelihood, Brier score"),
        ("_local", "accuracy re-weighted to each head's own clients' class mix"),
        ("_ens_bal / _ens_acc", "the averaged-softmax ensemble of all architectures (diagnostic)"),
        ("best_mean_acc", "the GeFL paper's metric: best test accuracy of each architecture over rounds, averaged (optimistic)"),
        ("fidelity_head / fidelity_tail", "accuracy of a referee classifier on generated features of common / rare classes"),
        ("cond_norm_*", "length of the generator's class rows (rare / common, or relative to the start)"),
        ("class_sep_F", "separation of class means relative to class spread (the gate statistic)"),
        ("*_time_s", "wall-clock seconds of a stage"),
        ("IF / alpha / K", "imbalance factor (0.01 = 100:1 long tail, 1.0 = balanced) / Dirichlet concentration of client shares / number of clients"),
        ("pfrac", "share of clients taking part in each round (blank = all)"),
        ("tail_perm", "rare classes drawn at random per seed instead of classes 6-9"),
        ("quick", "reduced-schedule viability pass, not a full run"),
        ("repeat_of_earlier_row", "this run id already appears earlier in the same experiment (a repeat, e.g. the reproducibility check)"),
    ]
    for g in gloss:
        gl.append(list(g))
    style_header(gl)
    set_widths(gl, [28, 110])
    for row in gl.iter_rows(min_row=2):
        for c in row:
            c.font = BODY
            c.alignment = Alignment(wrap_text=True, vertical="top")

    # ---- README
    n_exp = len({spec_of(f) for f in folders})
    lines = [
        ("All experiment results", TITLE),
        ("Generated by experiments/all_results_workbook.py from experiments/results and docs/research/math_checks. "
         "Re-run it after adding results; do not edit numbers by hand.", BODY),
        ("", BODY),
        (f"Logged runs (one row each in 'All runs'): {len(runs)}", BOLD),
        (f"Settings (one row each in 'Summary'): {len(groups)}", BOLD),
        (f"Experiments (specs): {n_exp}; result folders: {len(folders)}", BOLD),
        (f"MovieLens federated matrix-factorisation runs: {len(e40)}", BOLD),
        (f"Values from other result files (D01, T01, T02, CPU checks): {ot.max_row - 1}", BOLD),
        ("", BODY),
        ("How to read it", BOLD),
        ("Blank cell = that run did not measure that metric (metrics were added over time: e.g. readouts from K16, "
         "precision from 22 Oct 2026).", BODY),
        ("Accuracies, recalls, F1, precision and calibration error are shown as percentages; NLL, Brier and times as numbers.", BODY),
        ("'Summary' means are AVERAGEIF formulas over 'All runs' (blanks ignored); SD columns use the same rows.", BODY),
        ("'Experiments' also lists every planned Kaggle job whose results never came back, marked as such.", BODY),
        ("Not included: per-round training curves (the 'history' field), per-client class-count matrices and the nested "
         "instrumentation records of K25; they remain in the runs.jsonl files.", BODY),
    ]
    for t, f in lines:
        readme.append([t])
        readme.cell(row=readme.max_row, column=1).font = f
    readme.column_dimensions["A"].width = 140
    for row in readme.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT}: {len(runs)} runs, {len(groups)} settings, {len(folders)} folders, {len(e40)} E40 runs, "
          f"{ot.max_row - 1} other values, {pc.max_row - 1} per-class rows, {len(ordered)} metric columns")


if __name__ == "__main__":
    main()
